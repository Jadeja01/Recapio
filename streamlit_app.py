"""
AI-Powered Meeting Minutes Generator - Streamlit version
Built to fit inside Streamlit Community Cloud's free tier (~1GB RAM).

Key differences from the Gradio/HF-Spaces version:
- Uses Whisper "tiny" (much smaller than "base") for transcription.
- Uses a small distilled summarizer (sshleifer/distilbart-cnn-6-6) instead
  of full BART.
- Models are loaded ONE AT A TIME and explicitly deleted + garbage
  collected right after use, so peak memory never has more than one
  model resident at once. This trades speed for memory headroom.
- Real speaker diarization (pyannote) is OPTIONAL and opt-in via a
  checkbox, since it's the heaviest component. By default, a simple
  pause-based heuristic is used instead (see simple_speaker_split).
  This is NOT true diarization - it just alternates a label whenever
  there's a long pause between segments. It cannot recognize that a
  speaker who left and returned is the same person.

Deploy on https://share.streamlit.io for free:
  1. Push this file (as streamlit_app.py), requirements.txt and
     packages.txt to a public GitHub repo.
  2. Go to share.streamlit.io -> New app -> pick the repo/branch.
  3. Set "Main file path" to streamlit_app.py -> Deploy.
"""

import os
import re
import gc
import tempfile
from datetime import datetime
from collections import defaultdict

import streamlit as st
from docx import Document

st.set_page_config(page_title="AI Meeting Minutes Generator", page_icon="🎙️")

ACTION_PATTERNS = [
    r"\bwill\b", r"\bneeds? to\b", r"\bshould\b", r"\bgoing to\b",
    r"\bmust\b", r"\bplease\b", r"\baction item\b",
    r"\bby (monday|tuesday|wednesday|thursday|friday|saturday|sunday|\d{1,2}\s?(am|pm)?)\b",
]
ACTION_REGEX = re.compile("|".join(ACTION_PATTERNS), re.IGNORECASE)


# --------------------------------------------------------------------
# Pipeline stages - each one imports its own heavy library lazily and
# frees the model from memory as soon as it's done.
# --------------------------------------------------------------------

def transcribe_audio(audio_path):
    import whisper
    model = whisper.load_model("tiny")
    result = model.transcribe(audio_path, verbose=False)
    del model
    gc.collect()
    return result["segments"], result["text"]


def simple_speaker_split(segments, pause_threshold=1.2):
    """Lightweight, no-model speaker-turn heuristic (see module docstring).
    pause_threshold: seconds of silence between segments before flipping
    the speaker label. Lower = more sensitive to short pauses."""
    labeled = []
    current_speaker = 1
    prev_end = None
    for seg in segments:
        if prev_end is not None and (seg["start"] - prev_end) > pause_threshold:
            current_speaker = 2 if current_speaker == 1 else 1
        labeled.append({
            "speaker": f"Speaker {current_speaker}",
            "text": seg["text"].strip(),
            "start": seg["start"],
            "end": seg["end"],
        })
        prev_end = seg["end"]
    return labeled


def real_diarization(audio_path, hf_token):
    """True diarization via pyannote. Heavier - only run if user opts in."""
    import torchaudio
    # Newer torchaudio releases removed the old multi-backend dispatch
    # API (list_audio_backends / set_audio_backend). Some pyannote.audio
    # versions still call list_audio_backends() internally, which raises
    # AttributeError on those newer torchaudio builds. Shim it back in
    # as a harmless no-op so pyannote's internal check doesn't crash.
    if not hasattr(torchaudio, "list_audio_backends"):
        # Must be non-empty: pyannote/torchaudio-adjacent code indexes
        # into this list (e.g. backends[0]) to pick a default backend.
        # "soundfile" is the modern default backend name.
        torchaudio.list_audio_backends = lambda: ["soundfile"]

    from pyannote.audio import Pipeline
    pipeline = Pipeline.from_pretrained(
        "pyannote/speaker-diarization-3.1", use_auth_token=hf_token
    )
    diarization = pipeline(audio_path)
    speaker_segments = [
        {"start": t.start, "end": t.end, "speaker": s}
        for t, _, s in diarization.itertracks(yield_label=True)
    ]
    del pipeline
    gc.collect()
    return speaker_segments


def attach_real_speakers(segments, speaker_segments):
    def speaker_at(t):
        for s in speaker_segments:
            if s["start"] <= t <= s["end"]:
                return s["speaker"]
        return "Unknown"

    return [
        {
            "speaker": speaker_at((seg["start"] + seg["end"]) / 2),
            "text": seg["text"].strip(),
            "start": seg["start"],
            "end": seg["end"],
        }
        for seg in segments
    ]


def summarize_text(full_text):
    # Load the model/tokenizer directly instead of using the pipeline()
    # task-name shorthand ("summarization"). Some transformers releases
    # have changed or dropped that task-registry lookup, causing
    # KeyError: Unknown task summarization even though the model itself
    # works fine. Direct loading sidesteps that entirely.
    import torch
    from transformers import AutoTokenizer, AutoModelForSeq2SeqLM

    model_name = "sshleifer/distilbart-cnn-6-6"
    tokenizer = AutoTokenizer.from_pretrained(model_name)
    model = AutoModelForSeq2SeqLM.from_pretrained(model_name)
    model.eval()

    def chunk_text(text, max_words=500):
        words = text.split()
        for i in range(0, len(words), max_words):
            yield " ".join(words[i:i + max_words])

    parts = []
    with torch.no_grad():
        for chunk in chunk_text(full_text):
            inputs = tokenizer(
                chunk, return_tensors="pt", truncation=True, max_length=1024
            )
            summary_ids = model.generate(
                **inputs,
                max_length=100,
                min_length=20,
                num_beams=4,
                do_sample=False,
            )
            summary = tokenizer.decode(summary_ids[0], skip_special_tokens=True)
            parts.append(summary)

    del model
    del tokenizer
    gc.collect()
    return " ".join(parts)


def extract_action_items_and_dates(labeled_segments):
    """Single spaCy pass that returns two things:
    - action_items: lines matching ACTION_REGEX, with owner + due date
    - key_dates: every DATE entity mentioned anywhere in the transcript,
      with the speaker who said it and the sentence it appeared in
      (used for the Agenda / Key Dates section)."""
    import spacy
    nlp = spacy.load("en_core_web_sm")

    action_items = []
    key_dates = []

    for seg in labeled_segments:
        text = seg["text"]
        if not text:
            continue
        doc = nlp(text)
        dates_in_seg = [e.text for e in doc.ents if e.label_ == "DATE"]

        if ACTION_REGEX.search(text):
            people = [e.text for e in doc.ents if e.label_ == "PERSON"]
            owner = people[0] if people else seg["speaker"]
            due = dates_in_seg[0] if dates_in_seg else "Not specified"
            action_items.append({"task": text, "owner": owner, "due": due})

        for d in dates_in_seg:
            key_dates.append({
                "date": d,
                "speaker": seg["speaker"],
                "context": text,
            })

    del nlp
    gc.collect()
    return action_items, key_dates


def build_docx(summary, action_items, key_dates, labeled_segments):
    speaking_time = defaultdict(float)
    for seg in labeled_segments:
        speaking_time[seg["speaker"]] += seg["end"] - seg["start"]

    doc = Document()
    doc.add_heading("Meeting Minutes (Auto-Generated)", level=1)
    doc.add_paragraph(f"Date: {datetime.now().strftime('%d %b %Y')}")

    doc.add_heading("Participants", level=2)
    for speaker, secs in speaking_time.items():
        doc.add_paragraph(f"{speaker} - spoke {int(secs // 60)}m {int(secs % 60)}s", style="List Bullet")

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(summary)

    if key_dates:
        doc.add_heading("Key Dates Mentioned (Agenda)", level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Date", "Speaker", "Context"
        for d in key_dates:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = d["date"], d["speaker"], d["context"]

    doc.add_heading("Action Items", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Task", "Owner", "Due"
    for a in action_items:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = a["task"], a["owner"], a["due"]

    tmp_path = os.path.join(tempfile.gettempdir(), "Meeting_Minutes.docx")
    doc.save(tmp_path)
    return tmp_path


# --------------------------------------------------------------------
# UI
# --------------------------------------------------------------------
st.title("🎙️ AI-Powered Meeting Minutes Generator")
st.caption(
    "Lightweight build for free hosting (~1GB RAM): Whisper-tiny + a small "
    "distilled summarizer + rule-based action-item extraction. Models load "
    "one at a time and are released from memory immediately after use."
)

uploaded_file = st.file_uploader("Upload meeting audio (.wav / .mp3 / .m4a)", type=["wav", "mp3", "m4a"])

use_real_diarization = st.checkbox(
    "Attempt real speaker diarization (pyannote) - heavier, may fail on free hosting",
    value=False,
)
hf_token = None
pause_threshold = 1.2
if use_real_diarization:
    hf_token = st.text_input(
        "Hugging Face token (read access; needed for pyannote)", type="password"
    )
    st.caption(
        "Accept the model terms first at huggingface.co/pyannote/speaker-diarization-3.1 "
        "and huggingface.co/pyannote/segmentation-3.0, then generate a token at "
        "huggingface.co/settings/tokens."
    )
else:
    pause_threshold = st.slider(
        "Speaker-turn pause sensitivity (seconds)",
        min_value=0.3, max_value=3.0, value=1.2, step=0.1,
        help=(
            "Simple speaker-turn detection flips the speaker label whenever "
            "the gap between segments exceeds this value. Lower it if your "
            "audio has short pauses between speakers; raise it if one "
            "speaker's natural pauses are wrongly splitting them into two."
        ),
    )

if uploaded_file and st.button("Generate Minutes", type="primary"):
    suffix = os.path.splitext(uploaded_file.name)[1]
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(uploaded_file.read())
        audio_path = tmp.name

    try:
        with st.spinner("Transcribing audio (Whisper-tiny)..."):
            segments, full_text = transcribe_audio(audio_path)

        with st.spinner("Identifying speakers..."):
            if use_real_diarization and hf_token:
                try:
                    speaker_segments = real_diarization(audio_path, hf_token)
                    labeled_segments = attach_real_speakers(segments, speaker_segments)
                except Exception as e:
                    import traceback
                    st.warning(
                        f"Real diarization failed ({e}); falling back to simple speaker-turn detection."
                    )
                    with st.expander("Show full error details"):
                        st.code(traceback.format_exc())
                    labeled_segments = simple_speaker_split(segments, pause_threshold=pause_threshold)
            else:
                labeled_segments = simple_speaker_split(segments, pause_threshold=pause_threshold)

        with st.spinner("Summarizing discussion..."):
            summary = summarize_text(full_text)

        with st.spinner("Extracting action items and key dates..."):
            action_items, key_dates = extract_action_items_and_dates(labeled_segments)

        with st.spinner("Building minutes document..."):
            docx_path = build_docx(summary, action_items, key_dates, labeled_segments)
    finally:
        if os.path.exists(audio_path):
            os.remove(audio_path)
        gc.collect()

    st.subheader("Transcript (with speaker turns)")
    st.text("\n".join(f"[{s['speaker']}] {s['text']}" for s in labeled_segments))

    st.subheader("AI-Generated Summary")
    st.write(summary)

    st.subheader("Key Dates Mentioned (Agenda)")
    if key_dates:
        st.table(key_dates)
    else:
        st.write("No specific dates detected in the discussion.")

    st.subheader("Action Items")
    if action_items:
        st.table(action_items)
    else:
        st.write("No clear action items detected.")

    with open(docx_path, "rb") as f:
        st.download_button(
            "📄 Download Meeting Minutes (.docx)",
            f,
            file_name="Meeting_Minutes.docx",
        )